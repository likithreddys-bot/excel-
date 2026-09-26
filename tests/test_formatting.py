"""Phase 5: formatted Excel output - highlights, number formats, charts, default polish."""
import io

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

import app
import engine
from planner import make_plan
from test_planner import make_bank_df


@pytest.fixture(scope="module")
def sheets():
    return {"S": make_bank_df(500)}


def plan(sheets, cmd):
    p = make_plan(sheets, cmd)
    assert p.clarification_question is None, f"{cmd!r} -> asked: {p.clarification_question}"
    return p


def export(sheets, cmd):
    """Run the command, then write the Excel file with its formatting; return (data, workbook)."""
    p = plan(sheets, cmd)
    data = engine.apply_plan(sheets, p)
    formats = [s for s in p.steps if s.op in ("highlight", "number_format", "chart")]
    wb = openpyxl.load_workbook(io.BytesIO(engine.to_xlsx(data, formats)))
    return data, wb


def fill(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type else None


def only(p):
    (step,) = p.steps
    return step


# ---------- default polish ----------

def test_every_download_has_bold_frozen_header_and_widths(sheets):
    _, wb = export(sheets, "sort by amount")
    ws = wb["S"]
    assert ws.freeze_panes == "A2" and ws["A1"].font.b
    assert ws.column_dimensions["C"].width > 10  # description is wide


def test_formatting_does_not_change_data(sheets):
    p = plan(sheets, "highlight rows where amount > 100000 and show amount in rupees")
    assert engine.apply_plan(sheets, p)["S"].equals(sheets["S"])


# ---------- highlights ----------

def test_highlight_rows_in_red(sheets):
    df = sheets["S"]
    data, wb = export(sheets, "highlight rows where amount > 1 lakh in red")
    step = only(plan(sheets, "highlight rows where amount > 1 lakh in red"))
    assert (step.column, step.color) == (None, "F4CCCC")
    ws = wb["S"]
    for i in range(20):
        expected = "F4CCCC" if df.amount.iloc[i] > 100000 else None
        assert fill(ws.cell(i + 2, 1)) == expected and fill(ws.cell(i + 2, 8)) == expected


def test_highlight_cells_of_one_column(sheets):
    step = only(plan(sheets, "highlight amount above 150000"))
    assert step.column == "amount"
    _, wb = export(sheets, "highlight amount above 150000")
    i = int((sheets["S"].amount > 150000).to_numpy().nonzero()[0][0]) + 2
    assert fill(wb["S"].cell(i, 6)) == "FFF2CC" and fill(wb["S"].cell(i, 1)) is None


def test_highlight_bare_value_colours_rows(sheets):
    step = only(plan(sheets, "highlight debits in green"))
    assert (step.column, step.when.conditions[0].value) == (None, "DEBIT")


def test_highlight_duplicates_never_removes_rows(sheets):
    p = plan(sheets, "highlight duplicates in description")
    assert only(p).op == "highlight" and only(p).duplicates_in == "description"
    assert len(engine.apply_plan(sheets, p)["S"]) == len(sheets["S"])


def test_highlight_blanks(sheets):
    step = only(plan(sheets, "highlight blanks in branch"))
    assert (step.column, step.when.conditions[0].operator) == ("branch", "is_empty")


def test_highlight_survives_later_filter(sheets):
    data, wb = export(sheets, "highlight rows where amount > 150000 in red, then keep only debits")
    assert len(data["S"]) == (sheets["S"].txn_type == "DEBIT").sum()
    ws = wb["S"]
    amounts = [ws.cell(r, 6).value for r in range(2, ws.max_row + 1)]
    assert all((fill(ws.cell(r, 1)) == "F4CCCC") == (a > 150000) for r, a in zip(range(2, ws.max_row + 1), amounts))


def test_preview_shows_highlights(sheets):
    p = plan(sheets, "highlight rows where amount > 1 lakh in red")
    styles = engine.preview_styles(sheets["S"], p.steps, 50)
    assert [s["row"] == "F4CCCC" for s in styles] == (sheets["S"].amount.head(50) > 100000).tolist()


def test_highlight_note_counts_rows(sheets):
    notes = []
    engine.apply_plan(sheets, plan(sheets, "highlight rows where amount > 1 lakh"), notes=notes)
    assert f"{(sheets['S'].amount > 100000).sum():,} of 500 rows will be highlighted." in notes


# ---------- number formats ----------

def test_rupees_with_indian_grouping(sheets):
    _, wb = export(sheets, "show amount in rupees")
    assert "₹" in wb["S"].cell(2, 6).number_format and "##\\,##" in wb["S"].cell(2, 6).number_format


def test_highlighted_cells_keep_number_format(sheets):
    _, wb = export(sheets, "highlight rows where amount > 100000 in red and show amount in rupees")
    i = int((sheets["S"].amount > 100000).to_numpy().nonzero()[0][0]) + 2
    assert fill(wb["S"].cell(i, 6)) == "F4CCCC" and "₹" in wb["S"].cell(i, 6).number_format


@pytest.mark.parametrize("cmd, fmt", [
    ("format balance with commas", "#,##0.00"),
    ("show balance with 0 decimals", "#,##0"),
    ("show balance as whole numbers", "#,##0"),
    ("show balance to 1 decimal", "#,##0.0"),
])
def test_number_formats(sheets, cmd, fmt):
    _, wb = export(sheets, cmd)
    assert wb["S"].cell(2, 7).number_format == fmt


def test_date_format_converts_text_dates_visibly(sheets):
    p = plan(sheets, "show txn date as dd-mmm-yyyy")
    assert [s.op for s in p.steps] == ["convert", "number_format"]
    _, wb = export(sheets, "show txn date as dd-mmm-yyyy")
    assert wb["S"].cell(2, 2).number_format == "DD-MMM-YYYY"
    assert wb["S"].cell(2, 2).is_date


def test_percent_format_after_group(sheets):
    grouped = engine.apply_plan(sheets, plan(sheets, "total amount by category with % of total"))
    _, wb = export(grouped, "show % of total sum_amount as percent")
    assert wb["S"].cell(2, 3).number_format == '0.00"%"'


def test_number_format_on_text_column_stops(sheets):
    with pytest.raises(engine.PlanError, match="convert"):
        engine.apply_plan(sheets, plan(sheets, "show description in rupees"))


# ---------- charts ----------

def test_bar_chart_on_its_own_sheet(sheets):
    df = sheets["S"]
    data, wb = export(sheets, "bar chart of total amount by category")
    assert wb.sheetnames == ["S", "Chart 1"]
    ws = wb["Chart 1"]
    table = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert table == df.groupby("category").amount.sum().to_dict()
    assert len(ws._charts) == 1
    assert data["S"].equals(df)  # the data itself is not grouped


@pytest.mark.parametrize("cmd, kind, func, y, x, part", [
    ("line chart of amount by month", "line", "sum", "amount", "txn_date", "month"),
    ("pie chart of count by txn type", "pie", "count", None, "txn_type", None),
    ("horizontal bar chart of average amount by branch", "bar", "mean", "amount", "branch", None),
    ("plot total balance by year", "column", "sum", "balance", "txn_date", "year"),
])
def test_chart_variants(sheets, cmd, kind, func, y, x, part):
    s = only(plan(sheets, cmd))
    assert (s.kind, s.func, s.y, s.x, s.x_part) == (kind, func, y, x, part)


def test_chart_with_its_own_filter(sheets):
    df = sheets["S"]
    _, wb = export(sheets, "chart debits amount by category")
    ws = wb["Chart 1"]
    table = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert table == df[df.txn_type == "DEBIT"].groupby("category").amount.sum().to_dict()


def test_chart_of_it_after_grouping(sheets):
    p = plan(sheets, "total amount by category and add a pie chart of it by category")
    assert p.steps[-1].y == "sum_amount"


def test_chart_with_too_many_bars_asks(sheets):
    p = make_plan(sheets, "chart amount by txn id")
    assert p.steps == [] and "different values" in p.clarification_question


# ---------- through the app ----------

def test_app_download_has_formatting_and_keeps_summary_sheet_plain():
    main = pd.DataFrame({"pan": list("ABCD"), "result_code": [101, 109, 101, 109]})
    summary = pd.DataFrame({"result_code": [101, 109], "Count": [2, 2]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as w:
        main.to_excel(w, sheet_name="Output", index=False)
        summary.to_excel(w, sheet_name="Summary", index=False)
    c = TestClient(app.app)
    r = c.post("/api/upload", files={"file": ("itr.xlsx", buf.getvalue())}).json()
    sid = r["session_id"]
    plan_r = c.post("/api/plan", json={"session_id": sid, "message": "highlight rows where result code is 101 in green"}).json()
    assert "2 of 4 rows will be highlighted." in plan_r["notes"]
    st = c.post("/api/execute", json={"session_id": sid}).json()
    assert [s["row"] for s in st["sheets"][0]["styles"]] == ["D9EAD3", None, "D9EAD3", None]
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/download/{sid}").content))
    assert fill(wb["Output"]["A2"]) == "D9EAD3" and fill(wb["Output"]["A3"]) is None
    assert fill(wb["Summary"]["A2"]) is None  # the unchanged summary sheet is not highlighted
    c.post("/api/undo", json={"session_id": sid})
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/download/{sid}").content))
    assert fill(wb["Output"]["A2"]) is None  # undo removes the formatting too
